from pptx import Presentation

prs = Presentation()
methods_info = [
    ("random()", "Возвращает случайное число в диапазоне [0.0, 1.0)"),
    ("randint(a, b)", "Возвращает случайное целое число N, такое что a <= N <= b"),
    ("choice(seq)", "Возвращает случайный элемент из непустой последовательности seq"),
    ("shuffle(seq)", "Перемешивает последовательность seq на месте"),
    ("uniform(a, b)", "Возвращает случайное число с плавающей точкой в диапазоне [a, b]"),
    ("sample(population, k)", "Возвращает список из k уникальных элементов из population"),
]

for method, description in methods_info:
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    content = slide.shapes.placeholders[1]
    title.text = method
    content.text = description
    for paragraph in content.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.name = 'Courier New'
            run.font.size = 160000

output_path = "res.pptx"
prs.save(output_path)
