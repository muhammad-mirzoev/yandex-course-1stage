import openpyxl


def export_check(text):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    lines = text.strip().split('\n')
    row = 1
    for line in lines:
        product, price, quantity = line.split('\t')
        ws.cell(row=row, column=1, value=product)
        ws.cell(row=row, column=2, value=float(price))
        ws.cell(row=row, column=3, value=int(quantity))
        ws.cell(row=row, column=4, value=f"=B{row}*C{row}")

        row += 1

    ws.cell(row=row, column=1, value="Итого")
    ws.cell(row=row, column=4, value=f"=SUM(D1:D{row - 1})")
    wb.save("res.xlsx")


# Пример использования:
text = '''товар 1\t100\t5
товар 2\t200\t6
товар 3\t7\t500'''

export_check(text)
