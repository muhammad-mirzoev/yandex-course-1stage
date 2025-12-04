import openpyxl
from openpyxl.chart import PieChart, Reference

data = [
    ("Питание", 1200),
    ("Развлечения", 1500),
    ("Учеба", 300),
    ("Лечение", 100),
    ("Прочее", 670)
]

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Данные"

for i, (param, value) in enumerate(data, start=1):
    ws[f"A{i}"] = param
    ws[f"B{i}"] = value

chart = PieChart()
data_reference = Reference(ws, min_col=2, min_row=1, max_col=2, max_row=len(data))
categories_reference = Reference(ws, min_col=1, min_row=1, max_row=len(data))
chart.add_data(data_reference, titles_from_data=True)
chart.set_categories(categories_reference)

ws.add_chart(chart, "C3")
wb.save("res.xlsx")
