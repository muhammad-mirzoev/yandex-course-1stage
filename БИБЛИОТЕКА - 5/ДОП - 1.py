from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
import sys

data = list(map(lambda x: x.replace('\n', ''), sys.stdin.readlines()))
doc = Document()
z1 = doc.add_heading('Blood test', 0).alignment = 1

table = doc.add_table(len(data) + 1, 3)
table.style = 'Table Grid'
head_cells = table.rows[0].cells
for i, item in enumerate(['indicator', 'norm', 'value']):
    p = head_cells[i].paragraphs[0]
    p.add_run(item).bold = True
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
for i, item in enumerate(data):
    hdr_cells = table.rows[i + 1].cells
    hdr_cells[0].text = item

doc.save('analysis.docx')
