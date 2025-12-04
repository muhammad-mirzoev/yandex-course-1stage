from docx import Document
import sys

doc = Document()

spaces = {}
for line in list(map(str.rstrip, sys.stdin)):
    if len(line) - len(line.strip()) == 0:
        current_0 = line
        spaces[current_0] = {}
    elif len(line) - len(line.strip()) == 2:
        current_2 = line.strip()
        spaces[current_0][current_2] = {}
    elif len(line) - len(line.strip()) == 4:
        current_4 = line.strip()
        spaces[current_0][current_2][current_4] = []
    else:
        current_6 = line.strip()
        spaces[current_0][current_2][current_4].append(current_6)

count = 0
for w in spaces:
    if count == 0:
        doc.add_heading(w, level=0)
    doc.add_paragraph(w, style='List Bullet')
    for x in spaces[w]:
        doc.add_paragraph(x, style='List Bullet 2')
        for y in spaces[w][x]:
            doc.add_paragraph(y, style='List Bullet 3')
            for z in spaces[w][x][y]:
                doc.add_paragraph(z, style='List Bullet 4')
    count = 1

doc.save("systematization.docx")
